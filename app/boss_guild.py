from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from .extract import read_scores_from_image, merge_similar_names
from .compute import build_boss_guild_summary, generate_boss_guild_graphs

# ---------- CONFIG ----------
BASE_DIR = Path("folders")
BOSS_ROOT = BASE_DIR / "boss_guild"

OUTPUT_ROOT = Path("output")
BOSS_OUT_ROOT = OUTPUT_ROOT / "boss_guild"
BOSS_EXCEL_DIR = BOSS_OUT_ROOT / "excel"
BOSS_JSON_DIR = BOSS_OUT_ROOT / "json"

# สร้างโฟลเดอร์ที่ต้องใช้
for p in [OUTPUT_ROOT, BOSS_OUT_ROOT, BOSS_EXCEL_DIR, BOSS_JSON_DIR]:
    p.mkdir(parents=True, exist_ok=True)


def process_boss_guild():
    """อ่านทุกไฟล์ใน folders/boss_guild/* แล้วสรุปเป็น Excel + กราฟ"""

    all_rows = []

    if not BOSS_ROOT.exists():
        print("!! ไม่พบโฟลเดอร์", BOSS_ROOT)
        return

    # เดินทุกโฟลเดอร์ย่อยใน boss_guild (เช่น karma, kyle, tao, yon)
    for boss_folder in BOSS_ROOT.iterdir():
        if not boss_folder.is_dir():
            continue

        boss_name = boss_folder.name  # karma / kyle / tao / yon

        for pattern in ("*.png", "*.jpg", "*.jpeg"):
            for img_path in boss_folder.glob(pattern):
                print("อ่านรูป (boss):", img_path)
                all_rows.extend(
                    read_scores_from_image(img_path, boss_name, mode="boss")
                )

    if not all_rows:
        print("ไม่พบข้อมูลจากรูปใน boss_guild เลย")
        return

    df_raw = pd.DataFrame(all_rows)

    # รวมชื่อที่สะกดคล้ายกันให้ใช้ชื่อเดียวกัน
    df_raw["Member"] = merge_similar_names(df_raw["Member"], threshold=85)

    # -------------------- เก็บ raw OCR เป็น JSON --------------------
    run_date = datetime.today().strftime("%Y-%m-%d")
    raw_json_path = BOSS_JSON_DIR / f"boss_guild_raw_{run_date}.json"
    df_raw.to_json(raw_json_path, orient="records", force_ascii=False, indent=2)
    print("บันทึกข้อมูลดิบ boss_guild:", raw_json_path)

    # ---------- คำนวณสรุป + df สำหรับกราฟ ----------
    df_summary, df_graph = build_boss_guild_summary(df_raw)

    out_file = BOSS_EXCEL_DIR / "boss_guild_summary.xlsx"

    # ---------- เขียน Excel (แค่ชีทเดียว Summary) ----------
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        df_summary.to_excel(writer, sheet_name="Summary", index=False)

    # ---------- จัด format ----------
    wb = load_workbook(out_file)
    ws_sum = wb["Summary"]

    # จัด text ให้อยู่ตรงกลางทุกช่อง
    for row in ws_sum.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # ตั้งความกว้างคอลัมน์สำหรับ sheet Summary
    col_widths = {
        "A": 14,  # Members
        "B": 7,   # Score (M)
        "C": 10,  # Boss_name
        "D": 6,   # Time
        "E": 14,  # Total Dmg (Boss)
        "F": 14,  # Avg.Score (Boss)
        "G": 14,  # Total Dmg (All)
        "H": 14,  # Avg.Score (All)
        "I": 16,  # Contribution Rate per Boss
        "J": 16,  # Contribution Rate Total
        "K": 16,  # Boss Efficiency Index
        "L": 16,  # Guild Efficiency Index
        "M": 10,  # Best Boss
        "N": 10,  # Image
    }
    for col, width in col_widths.items():
        ws_sum.column_dimensions[col].width = width

    wb.save(out_file)

    print("\nบันทึกไฟล์ boss_guild เสร็จ:", out_file)
    print("ตัวอย่าง Summary:")
    print(df_summary.head())

    # ---------- สร้างกราฟ ----------
    generate_boss_guild_graphs(df_graph, OUTPUT_ROOT / "boss_guild_graphs")


if __name__ == "__main__":
    process_boss_guild()
