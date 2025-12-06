from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from .extract import read_scores_from_image, merge_similar_names
from .compute import build_castle_summary, generate_castle_graphs

# ---------- CONFIG ----------
BASE_DIR = Path("folders")
CASTLE_ROOT = BASE_DIR / "castles"

OUTPUT_ROOT = Path("output")
CASTLE_OUT_ROOT = OUTPUT_ROOT / "castle"
CASTLE_EXCEL_DIR = CASTLE_OUT_ROOT / "excel"
CASTLE_JSON_DIR = CASTLE_OUT_ROOT / "json"

for p in [OUTPUT_ROOT, CASTLE_OUT_ROOT, CASTLE_EXCEL_DIR, CASTLE_JSON_DIR]:
    p.mkdir(parents=True, exist_ok=True)


def process_castle():
    """
    อ่านทุกไฟล์ใน folders/castles/* (เช่น phy_castle, mag_castle, ...) แล้วสรุปเป็น Excel + กราฟ
    """

    all_rows = []

    if not CASTLE_ROOT.exists():
        print("!! ไม่พบโฟลเดอร์", CASTLE_ROOT)
        return

    # เดินทุกโฟลเดอร์ย่อยใน castles (เช่น phy_castle, mag_castle)
    for castle_folder in CASTLE_ROOT.iterdir():
        if not castle_folder.is_dir():
            continue

        castle_name = castle_folder.name  # phy_castle / mag_castle / ...

        for pattern in ("*.png", "*.jpg", "*.jpeg"):
            for img_path in castle_folder.glob(pattern):
                print("อ่านรูป (castle):", img_path)
                all_rows.extend(
                    read_scores_from_image(img_path, castle_name, mode="castle")
                )

    if not all_rows:
        print("ไม่พบข้อมูลจากรูปใน castles เลย")
        return

    df_raw = pd.DataFrame(all_rows)

    # รวมชื่อที่สะกดคล้ายกันให้ใช้ชื่อเดียวกัน
    df_raw["Member"] = merge_similar_names(df_raw["Member"], threshold=85)

    # -------------------- เก็บ raw OCR เป็น JSON --------------------
    run_date = datetime.today().strftime("%Y-%m-%d")
    raw_json_path = CASTLE_JSON_DIR / f"castle_raw_{run_date}.json"
    df_raw.to_json(raw_json_path, orient="records", force_ascii=False, indent=2)
    print("บันทึกข้อมูลดิบ castle:", raw_json_path)

    # ---------- คำนวณสรุป + df สำหรับกราฟ ----------
    df_summary, df_graph = build_castle_summary(df_raw)

    out_file = CASTLE_EXCEL_DIR / "castle_summary.xlsx"

    # ---------- เขียน Excel ----------
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        df_summary.to_excel(writer, sheet_name="Summary", index=False)

    # ---------- จัด format ----------
    wb = load_workbook(out_file)
    ws_sum = wb["Summary"]

    for row in ws_sum.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = {
        "A": 14,  # Members
        "B": 9,   # Score (M)
        "C": 12,  # Castle
        "D": 16,  # Total Dmg (Castle)
        "E": 16,  # Total Dmg (All)
        "F": 16,  # Avg.Score (All)
        "G": 20,  # Contribution Rate per Castle
        "H": 20,  # Contribution Rate Total
        "I": 18,  # Boss Efficiency Index
        "J": 18,  # Guild Efficiency Index
        "K": 12,  # Best Castle
        "L": 10,  # Image
    }
    for col, width in col_widths.items():
        ws_sum.column_dimensions[col].width = width

    wb.save(out_file)

    print("\nบันทึกไฟล์ castle เสร็จ:", out_file)
    print("ตัวอย่าง Summary:")
    print(df_summary.head())

    # ---------- กราฟ ----------
    generate_castle_graphs(df_graph, OUTPUT_ROOT / "castle_graphs")


if __name__ == "__main__":
    process_castle()
