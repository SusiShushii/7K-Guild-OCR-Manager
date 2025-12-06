from pathlib import Path
import re
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

# ---------- CONFIG ----------
OUTPUT_ROOT = Path("output")
CASTLE_JSON_DIR = OUTPUT_ROOT / "castle" / "json"
CASTLE_OUTPUT_DIR = OUTPUT_ROOT / "castle" / "growth"
CASTLE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)         # สร้างโฟลเดอร์อัตโนมัติ

# castle_raw_YYYY-MM-DD.json
JSON_PATTERN = re.compile(r"castle_raw_(\d{4}-\d{2}-\d{2})\.json")


def find_snapshot_files() -> list[Path]:
    if not CASTLE_JSON_DIR.exists():
        print(f"!! ไม่พบโฟลเดอร์: {CASTLE_JSON_DIR}")
        return []

    files: list[tuple[datetime, Path]] = []
    for p in CASTLE_JSON_DIR.glob("castle_raw_*.json"):
        m = JSON_PATTERN.fullmatch(p.name)
        if not m:
            continue
        dt = datetime.strptime(m.group(1), "%Y-%m-%d")
        files.append((dt, p))

    files.sort(key=lambda x: x[0])
    return [p for _, p in files]


def load_snapshot(path: Path, snapshot_date: datetime) -> pd.DataFrame:
    df = pd.read_json(path, orient="records")

    required = {"Member", "Boss_name", "Score_str"}
    if not required.issubset(df.columns):
        print(f"ไฟล์ {path} ไม่มีคอลัมน์ {required} ครบ ข้ามไฟล์นี้")
        return pd.DataFrame(columns=["Date", "Member", "Castle",
                                     "Total_raw", "Total_M"])

    df["Score_num"] = (
        df["Score_str"]
        .astype(str)
        .str.replace(",", "", regex=False)
        .astype("int64")
    )

    agg = (
        df.groupby(["Member", "Boss_name"])["Score_num"]
        .sum()
        .reset_index()
        .rename(columns={"Score_num": "Total_raw",
                         "Boss_name": "Castle"})
    )

    agg["Total_M"] = (agg["Total_raw"] / 1_000_000).round(2)
    agg["Date"] = snapshot_date.date()
    agg = agg[["Date", "Member", "Castle", "Total_raw", "Total_M"]]
    return agg


def build_growth_table() -> pd.DataFrame:
    snap_files = find_snapshot_files()
    if not snap_files:
        print("ไม่พบ snapshot castle_raw_*.json เลย")
        return pd.DataFrame()

    all_dfs = []
    for path in snap_files:
        m = JSON_PATTERN.fullmatch(path.name)
        dt = datetime.strptime(m.group(1), "%Y-%m-%d")
        print(f"โหลด snapshot: {path.name}")
        df_snap = load_snapshot(path, dt)
        if not df_snap.empty:
            all_dfs.append(df_snap)

    if not all_dfs:
        print("โหลด snapshot castle ไม่ได้เลย")
        return pd.DataFrame()

    df_all = pd.concat(all_dfs, ignore_index=True)
    df_all = df_all.sort_values(["Member", "Castle", "Date"])

    def calc_diff(g: pd.DataFrame) -> pd.DataFrame:
        g = g.sort_values("Date")
        g["Diff_M"] = g["Total_M"].diff().round(2)
        prev = g["Total_M"].shift(1)
        g["Pct_change"] = np.where(
            prev > 0,
            (g["Diff_M"] / prev * 100).round(1),
            np.nan,
        )
        return g

    df_growth = (
        df_all.groupby(["Member", "Castle"], group_keys=False)
        .apply(calc_diff)
    )
    return df_growth


def build_latest_growth(df_growth: pd.DataFrame) -> pd.DataFrame:
    rows = []

    for (member, castle), g in df_growth.groupby(["Member", "Castle"]):
        g = g.sort_values("Date")
        first = g.iloc[0]
        last = g.iloc[-1]

        first_total = float(first["Total_M"])
        last_total = float(last["Total_M"])
        abs_growth = round(last_total - first_total, 2)
        if first_total > 0:
            pct = round(abs_growth / first_total * 100, 1)
        else:
            pct = np.nan

        rows.append(
            {
                "Member": member,
                "Castle": castle,
                "First_Date": first["Date"],
                "First_Total_M": first_total,
                "Last_Date": last["Date"],
                "Last_Total_M": last_total,
                "Abs_Growth_M": abs_growth,
                "Growth_pct": pct,
            }
        )

    df_latest = pd.DataFrame(rows)
    df_latest = df_latest.sort_values(
        ["Member", "Abs_Growth_M"], ascending=[True, False]
    )
    return df_latest


def auto_format_excel(path: Path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center",
                                           vertical="center")
        for col_cells in ws.columns:
            col = col_cells[0].column_letter
            max_len = 0
            for c in col_cells:
                val = "" if c.value is None else str(c.value)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col].width = max(10, min(max_len + 2, 25))
    wb.save(path)


def run_castle_growth():
    df_growth = build_growth_table()
    if df_growth.empty:
        return

    df_latest = build_latest_growth(df_growth)

    df_timeline = df_growth.rename(columns={
        "Total_raw": "Total Damage (raw)",
        "Total_M": "Total Damage (M)",
        "Diff_M": "Damage Difference (M)",      # (Total_M - Previous Total_M)
        "Pct_change": "Damage Change (%)",      # (Diff_M / Previous Total_M) * 100
    })

    out_path = CASTLE_OUTPUT_DIR / "castle_growth_per_castle.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df_timeline.to_excel(writer, sheet_name="Timeline", index=False)
        df_latest.to_excel(writer, sheet_name="LatestGrowth", index=False)

    auto_format_excel(out_path)

    print("\nบันทึกไฟล์ growth (Castle) แล้วที่:")
    print(out_path)
    print("\nตัวอย่าง Timeline:")
    print(df_timeline.head())
    print("\nตัวอย่าง LatestGrowth:")
    print(df_latest.head())

if __name__ == "__main__":
    run_castle_growth()
